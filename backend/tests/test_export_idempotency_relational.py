from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import func, select

from app.main import create_app
from app.models.schema import exports, idempotency_records
from app.models.store import StoredObject
from tests.ocr_fakes import FakeOcrAdapter, wait_for_job


def test_export_and_idempotency_use_relational_tables(tmp_path: Path) -> None:
    app = create_app(str(tmp_path), f"sqlite:///{(tmp_path / 'milestone6.db').as_posix()}", FakeOcrAdapter())
    client = TestClient(app)
    auth = client.post("/api/v1/auth/register", json={
        "name": "导出用户", "phone": "13800134444",
        "password": "Export123", "employee_no": "EXPORT-001",
    }).json()["data"]
    headers = {"Authorization": f"Bearer {auth['access_token']}"}
    image = BytesIO()
    Image.new("RGB", (600, 400), "white").save(image, format="PNG")
    content = image.getvalue()
    file_id = client.post("/api/v1/files/upload-sessions", headers=headers, json={
        "file_name": "export.png", "size_bytes": len(content), "media_type": "image/png",
    }).json()["data"]["file_id"]
    client.put(f"/api/v1/files/{file_id}/content", headers=headers, content=content)
    client.post(f"/api/v1/files/{file_id}/complete", headers=headers)
    job_id = client.post("/api/v1/ocr/jobs", headers=headers, json={
        "name": "export job", "file_id": file_id,
        "model_id": "mock", "model_version": "1.0.0",
    }).json()["data"]["id"]
    wait_for_job(client, job_id, headers)
    result_id = client.get(
        f"/api/v1/ocr/jobs/{job_id}/pages/1/results", headers=headers
    ).json()["data"]["items"][0]["id"]

    idempotent_headers = {**headers, "Idempotency-Key": "same-correction"}
    first = client.post(
        f"/api/v1/ocr/results/{result_id}/corrections", headers=idempotent_headers,
        json={"corrected_text": "XT-200", "base_revision": 0},
    )
    second = client.post(
        f"/api/v1/ocr/results/{result_id}/corrections", headers=idempotent_headers,
        json={"corrected_text": "XT-200", "base_revision": 0},
    )
    assert first.status_code == second.status_code == 201
    assert first.json()["data"] == second.json()["data"]

    export = client.post(
        f"/api/v1/ocr/jobs/{job_id}/exports", headers=headers,
        json={"format": "xlsx", "mode": "simple", "scope": "all_pages"},
    ).json()["data"]
    assert client.get(
        f"/api/v1/ocr/jobs/{job_id}/exports/{export['id']}", headers=headers
    ).status_code == 200
    download = client.get(
        f"/api/v1/ocr/jobs/{job_id}/exports/{export['id']}/download", headers=headers
    )
    assert download.status_code == 200
    worksheet = load_workbook(BytesIO(download.content)).active
    assert worksheet.title == "精简结果"
    assert [cell.value for cell in worksheet[1]] == ["编号", "端子排最终识别结果"]
    assert worksheet.cell(row=2, column=1).value == "01"
    assert worksheet.cell(row=2, column=2).value == "XT-200"

    with app.state.service.store.engine.connect() as connection:
        assert connection.scalar(select(func.count()).select_from(exports)) == 1
        assert connection.scalar(select(func.count()).select_from(idempotency_records)) == 1
        legacy = connection.scalar(
            select(func.count()).select_from(StoredObject).where(
                StoredObject.kind.in_(["export", "idempotency"])
            )
        )
        assert legacy == 0
