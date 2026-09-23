from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import update

from app.models.schema import results
from app.services.terminal import merge_rows
from tests.test_api import client, workflow, register_headers
from tests.ocr_fakes import wait_for_job


def test_additive_migration_upgrade_and_downgrade(tmp_path):
    import importlib.util
    from pathlib import Path
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    from sqlalchemy import create_engine, inspect
    migration_path = Path(__file__).parents[1] / 'migrations/versions/f280_terminal_tasks.py'
    spec = importlib.util.spec_from_file_location('terminal_migration', migration_path)
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    engine = create_engine(f"sqlite:///{tmp_path / 'migration.db'}")
    with engine.begin() as connection:
        with Operations.context(MigrationContext.configure(connection)):
            migration.upgrade()
            assert set(inspect(connection).get_table_names()) == {'terminal_tasks', 'terminal_sides', 'terminal_reviews', 'terminal_events'}
            assert 'fingerprint' in {c['name'] for c in inspect(connection).get_columns('terminal_reviews')}
            migration.downgrade()
            assert not inspect(connection).get_table_names()


def source(id, number, text=None):
    return {"id": id, "terminal_number": number, "text": text or id, "table_note": ""}


def test_matching_is_by_unique_number_and_preserves_every_source():
    left = [source("l1", "02"), source("l2", "01"), source("l3", ""), source("l4", "D"), source("l5", "D")]
    right = [source("r1", "01"), source("r2", "02"), source("r3", ""), source("r4", "D"), source("r5", "1")]
    rows = merge_rows({"left": left, "right": right}, [], {})
    assert len(rows) == 8
    assert {(r["left"]["id"], r["right"]["id"]) for r in rows if r["left"] and r["right"]} == {("l1", "r2"), ("l2", "r1")}
    assert len([r[s]["id"] for r in rows for s in ("left", "right") if r[s]]) == 10
    assert any(r["conflict"] for r in rows)


def test_manual_single_does_not_automatically_rejoin_and_stale_pair_is_ignored():
    sources = {"left": [source("l", "01")], "right": [source("r", "01")]}
    rows = merge_rows(sources, [{"left": "l", "right": None}], {})
    assert len(rows) == 2
    rows = merge_rows(sources, [{"left": "l", "right": "removed"}], {})
    assert len(rows) == 1 and rows[0]["right"]["id"] == "r"


def setup_pair(api):
    left, headers = workflow(api)
    right = api.post(f"/api/v1/ocr/jobs/{left}/rerun", headers=headers,
        json={"model_id": "mock", "model_version": "1.0.0"}).json()["data"]["id"]
    wait_for_job(api, right, headers)
    for job in (left, right):
        items = api.get(f"/api/v1/ocr/jobs/{job}/pages/1/results", headers=headers).json()["data"]["items"]
        for i, item in enumerate(items):
            response = api.patch(f"/api/v1/ocr/results/{item['id']}/table", headers=headers,
                json={"terminal_number": f"{i + 1:02}", "manual_confirmed": False, "table_note": "", "base_revision": 0})
            assert response.status_code == 200, response.text
    created = api.post("/api/v1/terminal-tasks", headers=headers, json={"name": "X1"})
    assert created.status_code == 200, created.text
    task = created.json()["data"]
    for side, job in (("left", left), ("right", right)):
        response = api.put(f"/api/v1/terminal-tasks/{task['id']}/sides/{side}", headers=headers,
            json={"job_id": job, "page_no": 1, "base_revision": task["revision"]})
        assert response.status_code == 200, response.text
        task = response.json()["data"]
    return task, headers


def version(task):
    return {"base_revision": task["revision"], "version": task["version"]}


def test_association_review_export_snapshot_and_source_changes(tmp_path):
    with client(tmp_path) as api:
        task, headers = setup_pair(api)
        path = f"/api/v1/terminal-tasks/{task['id']}"
        assert len(task["rows"]) == 3
        row = task["rows"][0]
        response = api.patch(path + "/rows", headers=headers, json=version(task) | {
            "rows": [{"id": row["id"], "confirmed": True, "note": "=literal"}]})
        assert response.status_code == 200, response.text
        task = response.json()["data"]
        assert task["rows"][0]["confirmed"]
        export = api.post(path + "/exports", headers=headers, json=version(task) | {"scope": "all"})
        assert export.status_code == 200, export.text
        url = export.json()["data"]["download_url"]
        book = load_workbook(BytesIO(api.get(url, headers=headers).content))
        assert list(next(book.active.values)) == ["序号", "端子编号", "左侧线缆号", "右侧线缆号", "人工确认", "备注"]
        assert book.active["F2"].value == "=literal" and book.active["F2"].data_type == "s"
        corrected = api.post(f"/api/v1/ocr/results/{row['left']['id']}/corrections", headers=headers,
            json={"corrected_text": "CHANGED", "base_revision": 0})
        assert corrected.status_code == 201
        stale = api.patch(path + "/rows", headers=headers, json=version(task) | {"rows": [{"id": row["id"], "confirmed": True, "note": ""}]})
        assert stale.status_code == 409
        fresh = api.get(path, headers=headers).json()["data"]
        assert not fresh["rows"][0]["confirmed"]
        assert fresh["rows"][0]["left"]["text"] == "CHANGED"
        snapshot = load_workbook(BytesIO(api.get(url, headers=headers).content))
        assert snapshot.active["C2"].value != "CHANGED"
        assert snapshot.active["E2"].value == "是"


def test_permissions_idempotency_and_invalid_associations(tmp_path):
    with client(tmp_path) as api:
        task, headers = setup_pair(api)
        other = register_headers(api, "13800138000", "OTHER")
        path = f"/api/v1/terminal-tasks/{task['id']}"
        assert api.get(path, headers=other).status_code == 404
        assert api.get('/api/v1/terminal-tasks', headers=other).json()['data']['items'] == []
        assert api.get(path).status_code == 401
        payload = {"job_id": task["sides"]["left"]["job_id"], "page_no": 1, "base_revision": task["revision"]}
        assert api.put(path + '/sides/right', headers=headers, json=payload).status_code == 409
        assert api.put(path + '/sides/left', headers=headers, json=payload | {"page_no": 99}).status_code == 422
        assert api.put(path + '/sides/left', headers=headers, json=payload | {"base_revision": 0}).status_code == 409
        assert api.put(path + '/sides/left', headers=other, json=payload).status_code == 404
        foreign = api.post('/api/v1/terminal-tasks', headers=other, json={"name": "Other"}).json()['data']
        assert api.put(f"/api/v1/terminal-tasks/{foreign['id']}/sides/left", headers=other, json=payload | {"base_revision": 0}).status_code == 404
        key = headers | {"Idempotency-Key": str(uuid4())}
        first = api.post('/api/v1/terminal-tasks', headers=key, json={"name": "Idempotent"}).json()['data']
        second = api.post('/api/v1/terminal-tasks', headers=key, json={"name": "Idempotent"}).json()['data']
        assert first['id'] == second['id']
        assert api.post('/api/v1/terminal-tasks', headers=key, json={"name": "Changed"}).status_code == 409


def test_manual_unlink_relink_conflicts_and_deleted_source(tmp_path):
    with client(tmp_path) as api:
        task, headers = setup_pair(api)
        path = f"/api/v1/terminal-tasks/{task['id']}"
        row = task['rows'][0]
        task = api.post(path + '/links', headers=headers, json=version(task) | {"left": row['left']['id'], "right": None}).json()['data']
        assert len(task['rows']) == 4
        task = api.post(path + '/links', headers=headers, json=version(task) | {"left": row['left']['id'], "right": row['right']['id']}).json()['data']
        assert len(task['rows']) == 3
        with api.app.state.service.store.engine.begin() as db:
            db.execute(update(results).where(results.c.id == task['rows'][1]['left']['id']).values(terminal_number='01'))
        task = api.get(path, headers=headers).json()['data']
        assert any(r['conflict'] for r in task['rows'])
        assert api.post(path + '/exports', headers=headers, json=version(task) | {"scope": "all"}).status_code == 409
        # Deleted sources are visibly unavailable, not silently exported as blanks.
        api.delete(f"/api/v1/ocr/jobs/{task['sides']['right']['job_id']}", headers=headers)
        task = api.get(path, headers=headers).json()['data']
        assert task['sides']['right']['status'] == 'unavailable'
